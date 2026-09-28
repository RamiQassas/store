import io
import openpyxl
from decimal import Decimal
from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings, Client
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from apps.catalog.models import Category, Product, ProductVariant

User = get_user_model()

TEST_CACHES = {
    "default": {
        "BACKEND": "django.core.cache.backends.locmem.LocMemCache",
        "LOCATION": "tests",
    }
}

@override_settings(SECURE_SSL_REDIRECT=False, CACHES=TEST_CACHES)
class ProductExcelImportTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(
            email="admin@example.com",
            password="AdminPassword123!",
            role=User.Role.SUPER_ADMIN
        )
        self.client = Client()
        self.client.force_login(self.admin)

    def test_download_excel_template(self):
        """Test that downloading the import template returns a valid Excel file with 2 sheets."""
        url = reverse("control_product_import_template")
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("product_import_template.xlsx", response["Content-Disposition"])

        # Parse downloaded workbook
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertIn("قالب استيراد المنتجات", wb.sheetnames)
        self.assertIn("تعليمات الاستيراد الهامة", wb.sheetnames)

        ws = wb["قالب استيراد المنتجات"]
        headers = [cell.value for cell in ws[1] if cell.value]
        self.assertIn("اسم المنتج *", headers)
        self.assertIn("نوع المنتج * (رقمي / مادي)", headers)
        self.assertIn("القسم / التصنيف *", headers)
        self.assertIn("السعر (سعر البيع) *", headers)

        # Check example rows
        types_in_sheet = [ws.cell(row=r, column=3).value for r in range(2, 4)]
        self.assertIn("رقمي", types_in_sheet)
        self.assertIn("مادي", types_in_sheet)

    def test_import_physical_and_digital_products(self):
        """Test importing products from Excel, verifying physical vs digital classification and inventory."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "قالب استيراد المنتجات"

        # Headers
        ws.append([
            "معرّف المنتج (ID)",
            "اسم المنتج *",
            "نوع المنتج * (رقمي / مادي)",
            "القسم / التصنيف *",
            "اسم الباقة / الخيار",
            "السعر (سعر البيع) *",
            "التكلفة",
            "سعر الجملة",
            "سعر VIP",
            "الكمية المتوفرة",
            "رمز التخزين (SKU)",
            "نشط (نعم / لا)",
            "طريقة التسليم (يدوي / تلقائي)",
            "الوصف",
            "تعليمات الاستخدام",
        ])

        # Row 1: Physical product
        ws.append([
            "",
            "شاحن سريع أنكر 65W (Anker Fast Charger)",
            "مادي",
            "إلكترونيات وشواحن",
            "شاحن 65 واط أسود",
            45.00,
            30.00,
            38.00,
            36.00,
            25,
            "ANKER-65W-BLK",
            "نعم",
            "يدوي",
            "شاحن جداري فائق السرعة يدعم الشحن السريع للهواتف والحواسيب",
            "توصيل سريع لباب المنزل",
        ])

        # Row 2: Digital product (Variant 1)
        ws.append([
            "",
            "شاهد VIP اشتراك رسمي (Shahid VIP)",
            "رقمي",
            "خدمات التلفزيون والبث",
            "اشتراك شهر واحد",
            9.99,
            7.50,
            8.50,
            8.00,
            0,
            "SHAHID-1M",
            "نعم",
            "تلقائي",
            "اشتراك رسمي لمشاهدة أقوى الأفلام والمسلسلات بجودة FHD",
            "يتم إرسال كود التفعيل فوراً",
        ])

        # Row 3: Digital product (Variant 2 for same product)
        ws.append([
            "",
            "شاهد VIP اشتراك رسمي (Shahid VIP)",
            "رقمي",
            "خدمات التلفزيون والبث",
            "اشتراك سنة كاملة",
            89.99,
            65.00,
            75.00,
            70.00,
            0,
            "SHAHID-1Y",
            "نعم",
            "تلقائي",
            "اشتراك رسمي لمشاهدة أقوى الأفلام والمسلسلات بجودة FHD",
            "يتم إرسال كود التفعيل فوراً",
        ])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        uploaded = SimpleUploadedFile(
            "test_products.xlsx",
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        url = reverse("control_product_import")
        response = self.client.post(url, {"file": uploaded})
        self.assertEqual(response.status_code, 302)

        # 1. Verify Physical Product
        p_phys = Product.objects.filter(name="شاحن سريع أنكر 65W (Anker Fast Charger)").first()
        self.assertIsNotNone(p_phys)
        self.assertEqual(p_phys.product_type, "physical")
        self.assertEqual(p_phys.category.name, "إلكترونيات وشواحن")
        self.assertTrue(p_phys.track_inventory)
        self.assertEqual(p_phys.quantity, 25)

        v_phys = p_phys.variants.first()
        self.assertIsNotNone(v_phys)
        self.assertEqual(v_phys.price, Decimal("45.00"))
        self.assertEqual(v_phys.cost, Decimal("30.00"))
        self.assertEqual(v_phys.delivery_type, "manual")
        self.assertEqual(v_phys.sku, "ANKER-65W-BLK")

        # 2. Verify Digital Product with multiple variants
        p_digi = Product.objects.filter(name="شاهد VIP اشتراك رسمي (Shahid VIP)").first()
        self.assertIsNotNone(p_digi)
        self.assertEqual(p_digi.product_type, "digital")
        self.assertEqual(p_digi.category.name, "خدمات التلفزيون والبث")
        self.assertEqual(p_digi.variants.count(), 2)

        v_1m = p_digi.variants.filter(name="اشتراك شهر واحد").first()
        self.assertIsNotNone(v_1m)
        self.assertEqual(v_1m.price, Decimal("9.99"))
        self.assertEqual(v_1m.delivery_type, "keys")

        v_1y = p_digi.variants.filter(name="اشتراك سنة كاملة").first()
        self.assertIsNotNone(v_1y)
        self.assertEqual(v_1y.price, Decimal("89.99"))
        self.assertEqual(v_1y.delivery_type, "keys")

    def test_update_existing_product_via_import(self):
        """Test updating an existing product's price, quantity, and name via Excel import."""
        cat = Category.objects.create(name="أجهزة")
        prod = Product.objects.create(
            name="ماوس ألعاب",
            product_type="physical",
            category=cat,
            quantity=10,
            track_inventory=True
        )
        ProductVariant.objects.create(
            product=prod,
            name="الأساسي",
            sku="MOUSE-001",
            price=Decimal("15.00")
        )

        # Create Excel with updated data using the product ID
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append([
            "ID", "اسم المنتج", "النوع", "القسم", "الباقة", "السعر", "الكمية", "SKU"
        ])
        ws.append([
            str(prod.id), "ماوس ألعاب لاسلكي مطور", "مادي", "أجهزة", "الأساسي", 25.00, 50, "MOUSE-001"
        ])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        uploaded = SimpleUploadedFile(
            "update_products.xlsx",
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        response = self.client.post(reverse("control_product_import"), {"file": uploaded})
        self.assertEqual(response.status_code, 302)

        prod.refresh_from_db()
        self.assertEqual(prod.name, "ماوس ألعاب لاسلكي مطور")
        self.assertEqual(prod.quantity, 50)
        self.assertEqual(prod.variants.first().price, Decimal("25.00"))
